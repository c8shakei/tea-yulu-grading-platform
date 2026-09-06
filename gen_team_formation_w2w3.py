#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
A 组（W2+W3）智能体团队组建与任务派发过程文档生成器
"""
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from datetime import datetime
import os

OUTPUT = r"C:\Users\wzd\Desktop\毕业设计\团队组建与任务派发_W2W3_A组.docx"

def set_default_font(doc, name="宋体", size=10.5):
    doc.styles["Normal"].font.name = name
    doc.styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    doc.styles["Normal"].font.size = Pt(size)

def add_heading_zh(doc, text, level=1):
    style = "Heading %d" % level
    p = doc.add_paragraph(style=style)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.font.name = "黑体" if level <= 2 else "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体" if level <= 2 else "宋体")
    run.font.bold = True
    run.font.size = Pt({1: 16, 2: 14}.get(level, 12))
    return p

def add_para(doc, text, bold=False, indent=True, size=10.5):
    p = doc.add_paragraph(style="Normal")
    if indent:
        p.paragraph_format.first_line_indent = Inches(0.3)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    run = p.add_run(text)
    run.font.bold = bold
    run.font.size = Pt(size)
    return p

def add_table(doc, rows, cols, data):
    table = doc.add_table(rows=rows, cols=cols)
    table.style = "Table Grid"
    for i, row_data in enumerate(data):
        row = table.rows[i]
        for j, cell_text in enumerate(row_data):
            cell = row.cells[j]
            cell.text = str(cell_text)
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    return table

def main():
    doc = Document()
    set_default_font(doc)

    # Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("A 组智能体团队组建与任务派发记录\n（W2 视觉模型 + W3 后端溯源）")
    run.font.name = "黑体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    run.font.size = Pt(20)
    run.font.bold = True

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run("项目：基于计算机视觉的农产品品质分级与溯源平台（恩施玉露茶）\n生成日期：2026-09-04")
    run.font.size = Pt(12)

    add_heading_zh(doc, "一、组建决策", level=1)
    add_para(doc, "用户（组长）已启用智能体团队管理功能，并选择方案 A：组建 W2+W3 并行小分队。选择理由：")
    add_para(doc, "• W2 视觉模型与 W3 后端溯源在接口契约层面可解耦推进；")
    add_para(doc, "• W3 的 FastAPI + SQLite + 哈希链模块不依赖 W2 的成品权重，可用 mock/stub 先行验证 API 与链式结构；")
    add_para(doc, "• 该方案在 30 天硬基线下风险最低，能最大限度并行压缩工期。")

    add_heading_zh(doc, "二、团队成员与职责", level=1)
    team = [
        ["角色", "智能体名称", "任务编号", "职责简述"],
        ["W2 视觉模型工程师", "w2-vision-engineer", "Task #16", "Colab GPU 训练 YOLOv8 检测主模型 + yolov8n-cls 对照模型；封装 detect(image) 推理接口"],
        ["W3 后端与溯源工程师", "w3-backend-engineer", "Task #17", "FastAPI + SQLite 后端；哈希链不可篡改验证；提供 OpenAPI 文档供 W4 对接"],
        ["组长（本对话）", "—", "—", "统筹、接口契约冻结、里程碑验收、问题裁决"],
    ]
    add_table(doc, len(team), 4, team)

    add_heading_zh(doc, "三、任务派发记录", level=1)
    add_heading_zh(doc, "3.1 W2 任务", level=2)
    w2_task = [
        ["字段", "内容"],
        ["任务名称", "启动 W2 视觉模型训练（M2）"],
        ["任务编号", "Task #16"],
        ["里程碑", "M2（D10≈2026-09-14）模型达标"],
        ["输入", "W1 交付的 data.yaml、datasets/tea_yulu、《数据说明.md》第 7 节 Colab 取数方案"],
        ["输出", "src/models/train_detect.py、src/models/train_cls.py、models/best_detect.pt、models/best_cls.pt、src/inference/detector.py、docs/W2_M2_*.md、W2_M2里程碑汇报_模型训练.docx"],
        ["关键约束", "训练在 Colab GPU；本地只验证 CPU 推理；接口契约②输出字段必须对齐；不得改类别/划分"],
    ]
    add_table(doc, len(w2_task), 2, w2_task)

    add_heading_zh(doc, "3.2 W3 任务", level=2)
    w3_task = [
        ["字段", "内容"],
        ["任务名称", "启动 W3 后端与溯源模块（M3）"],
        ["任务编号", "Task #17"],
        ["里程碑", "M3（D16≈2026-09-20）后端可用"],
        ["输入", "SRS 第4-5章接口契约、接口契约③④、W2 detect() 接口签名（先用 mock）"],
        ["输出", "src/backend/main.py、src/backend/routers/detect.py、src/backend/routers/trace.py、src/backend/db.py、src/hashchain/chain.py、src/backend/mock_detector.py、docs/W3_M3_*.md、W3_M3里程碑汇报_后端可用.docx"],
        ["关键约束", "统一响应 {code,message,data}；SQLite 单文件；哈希链必须验证不可篡改；接口字段与契约③一致"],
    ]
    add_table(doc, len(w3_task), 2, w3_task)

    add_heading_zh(doc, "四、执行与验收复核", level=1)
    add_para(doc, "两组智能体在 2026-09-04 15:27 启动后快速完成代码脚手架，组长随即进行复核。")
    add_para(doc, "W2 复核：运行 src/inference/detector.py，接口字段校验通过；当前本地无 best_detect.pt，脚本自动降级使用 yolov8n.pt 完成格式验证。真实训练待 Colab 导出权重后回填。")
    add_para(doc, "W3 复核：首次运行 scripts/test_m3.py 因环境缺失 python-multipart 失败；安装后复测 6 项全部 PASS。")

    check = [
        ["复核项", "结果", "备注"],
        ["W2 detector.py 接口字段", "PASS", "输出 {class_id, class_name, confidence, bbox}，bbox 归一化 [0,1]"],
        ["W2 train_detect.py/train_cls.py", "已产出", "待 Colab GPU 训练后回填指标"],
        ["W3 health / openapi / detect / trace POST / trace GET", "PASS", "统一响应 {code,message,data}"],
        ["W3 哈希链不可篡改验证", "PASS", "篡改 data_hash 后 verify_chain() 返回 False"],
        ["W3 mock_detector bbox 归一化", "已修复", "原 mock 返回像素值，已改为 [0,1] 与契约②一致"],
    ]
    add_table(doc, len(check), 3, check)

    add_heading_zh(doc, "五、问题与修复", level=1)
    add_heading_zh(doc, "5.1 python-multipart 缺失", level=2)
    add_para(doc, "现象：W3 后端使用 FastAPI UploadFile 依赖 multipart，本机 venv 未安装 python-multipart，导致 scripts/test_m3.py 导入失败。")
    add_para(doc, "修复：执行 pip install python-multipart（版本 0.0.32）后复测通过。")
    add_para(doc, "后续：环境配置清单.xlsx 应同步追加 python-multipart 项；W5 集成测试前再次检查。")

    add_heading_zh(doc, "5.2 mock_detector bbox 未归一化", level=2)
    add_para(doc, "现象：src/backend/mock_detector.py 返回的 bbox 为像素值 [50,150,400,400]，与接口契约②的归一化 [0,1] 不一致。")
    add_para(doc, "修复：改为返回 [0.05~0.25, 0.05~0.25, 0.70~0.95, 0.70~0.95] 范围内的归一化坐标。")
    add_para(doc, "影响：W4 前端按归一化值渲染时，mock 阶段不会出现缩放错位；W2 真实 detector 返回的同样为归一化值，接口一致。")

    add_heading_zh(doc, "六、当前状态与阻塞", level=1)
    status = [
        ["任务", "状态", "阻塞/待办"],
        ["W2 M2 模型训练", "进行中（代码完成）", "需上传数据集到 Google Drive；需在 Colab 完成训练并导出 best_detect.pt / best_cls.pt 到本地"],
        ["W3 M3 后端可用", "进行中（功能完成）", "等待 W2 真实权重，将 mock_detector 替换为 src.inference.detector.detect 并复测"],
    ]
    add_table(doc, len(status), 3, status)

    add_heading_zh(doc, "七、下一步", level=1)
    add_para(doc, "1. 用户将 D:\\BISHE_DATA\\datasets\\tea_yulu 上传至 Google Drive（推荐 /MyDrive/BISHE_DATA/datasets/tea_yulu）。")
    add_para(doc, "2. 在 Colab 运行 W2 的 train_detect.py 与 train_cls.py，导出权重和指标文件到本地 models/。")
    add_para(doc, "3. W3 将 mock_detector 替换为真实 detector，运行 scripts/test_m3.py 做回归验证。")
    add_para(doc, "4. W2/W3 分别按汇报模板提交 M2/M3 里程碑汇报。")
    add_para(doc, "5. 组长验收后启动 W4 前端工程师（依赖 W3 API 契约）。")

    add_heading_zh(doc, "八、过程文件清单", level=1)
    files = [
        ["类型", "路径", "说明"],
        ["团队组建记录", "团队组建与任务派发_W2W3_A组.docx", "本文档"],
        ["M1 批复", "M1验收批复书_W2启动令.docx", "批准 W2 启动的依据"],
        ["W2 代码", "src/models/*.py、src/inference/detector.py", "训练与推理封装"],
        ["W2 文档", "docs/W2_M2_训练说明.md、docs/W2_M2_对比实验.md、W2_M2里程碑汇报_模型训练.docx", "训练与对比说明"],
        ["W3 代码", "src/backend/*.py、src/backend/routers/*.py、src/hashchain/chain.py", "后端与哈希链"],
        ["W3 文档", "docs/W3_M3_API文档.md、docs/W3_M3_数据库说明.md、W3_M3里程碑汇报_后端可用.docx", "后端说明"],
        ["测试脚本", "scripts/test_m3.py", "W3 集成测试"],
        ["生成器", "gen_team_formation_w2w3.py", "本文档生成脚本"],
    ]
    add_table(doc, len(files), 3, files)

    doc.save(OUTPUT)
    print("OK saved:", OUTPUT)

if __name__ == "__main__":
    main()
