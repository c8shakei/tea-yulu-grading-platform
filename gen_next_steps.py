# -*- coding: utf-8 -*-
"""生成 W2/W3 A组后续操作步骤 docx"""
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

OUT = "C:\\Users\\wzd\\Desktop\\毕业设计\\W2W3_A组后续操作步骤.docx"

doc = Document()

# 样式
style = doc.styles['Normal']
style.font.name = 'Microsoft YaHei'
style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
style.font.size = Pt(10.5)

# 标题
title = doc.add_heading('W2/W3 A组后续操作步骤', level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# 引言
doc.add_paragraph('本步骤面向组长（用户）与 A 组两名智能体员工（W2、W3）。当前状态：W3 后端已跑通，W2 训练代码已就绪，缺真实权重。按以下顺序执行。')

def h2(t): doc.add_heading(t, level=2)
def h3(t): doc.add_heading(t, level=3)
def p(t, bold=False): doc.add_paragraph(t, style='List Bullet' if not bold else None)

def code_block(txt):
    p = doc.add_paragraph()
    run = p.add_run(txt)
    run.font.name = 'Consolas'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Consolas')
    run.font.size = Pt(9)
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.space_after = Pt(6)

h2('第一阶段：用户准备数据集（必须先做，约 5–15 分钟）')
h3('步骤 1.1 确认本地数据完整')
p('打开 PowerShell 或 CMD，执行以下命令检查 junction 与数据：')
code_block("""cd "C:\\Users\\wzd\\Desktop\\毕业设计"
ls datasets/tea_yulu/train/images | Measure-Object
ls datasets/tea_yulu/val/images | Measure-Object
ls datasets/tea_yulu/test/images | Measure-Object
""")
p('预期：train 约 4611 张、val 约 439 张、test 约 219 张。')

h3('步骤 1.2 把 D 盘数据打包上传到 Google Drive')
p('由于 Colab 看不到本机 D 盘，需要先把数据集上传到 Google Drive。推荐方式：')
p('方式 A（推荐）：直接压缩 D 盘数据集目录，通过浏览器上传到 Google Drive。')
code_block("""# PowerShell 中执行，生成压缩包
Compress-Archive -Path "D:\\BISHE_DATA\\datasets\\tea_yulu" `
  -DestinationPath "D:\\BISHE_DATA\\tea_yulu_dataset.zip" -Force
# 然后打开 https://drive.google.com，把 zip 拖到 Drive 根目录
""")
p('方式 B：把 datasets/tea_yulu（junction 后的项目路径）打包，结果一样。')

h3('步骤 1.3 打开 Colab 并挂载 Drive')
code_block("""# 在 Colab 新建 notebook，先跑这两行
from google.colab import drive
drive.mount('/content/drive')
""")

h2('第二阶段：W2 在 Colab 训练模型')
h3('步骤 2.1 解压数据集到 Colab')
code_block("""!unzip -q /content/drive/MyDrive/tea_yulu_dataset.zip -d /content/
!ls /content/tea_yulu
""")

h3('步骤 2.2 上传 W2 训练脚本')
p('把本地以下两个文件上传到 Colab（直接拖进文件区或用代码编辑器新建）：')
p('src/models/train_detect.py')
p('src/models/train_cls.py')
p('再上传 data.yaml 的便携版：D:\\BISHE_DATA\\datasets\\tea_yulu\\data.yaml（path 是 .）')

h3('步骤 2.3 安装依赖并运行训练')
code_block("""!pip install -q ultralytics

# 主检测模型
!python train_detect.py

# 对照分类模型
!python train_cls.py
""")

h3('步骤 2.4 导出 best.pt 到本地')
code_block("""# 训练完成后，把权重复制到 Drive，方便本机下载
!mkdir -p /content/drive/MyDrive/tea_yulu_outputs
!cp /content/tea_yulu/runs/detect/train/weights/best.pt /content/drive/MyDrive/tea_yulu_outputs/best_detect.pt
!cp /content/tea_yulu/runs/classify/train/weights/best.pt /content/drive/MyDrive/tea_yulu_outputs/best_cls.pt
!cp -r /content/tea_yulu/runs/detect/train /content/drive/MyDrive/tea_yulu_outputs/detect_train_results
""")
p('然后在本机下载到项目目录：')
code_block("""# PowerShell
cd "C:\\Users\\wzd\\Desktop\\毕业设计"
mkdir -Force models
# 从 Google Drive 网页下载 best_detect.pt / best_cls.pt，放到 models/ 目录
ls models
""")

h2('第三阶段：W3 替换 mock 并集成验证')
h3('步骤 3.1 验证 best_detect.pt 存在')
code_block("""# 在项目根目录 PowerShell
ls models/best_detect.pt
ls models/best_cls.pt
""")

h3('步骤 3.2 修改 W3 后端使用真实检测器')
p('编辑 src/backend/main.py，把 mock_detector 替换为真实检测器：')
code_block("""# src/backend/main.py
# 原来：
# from backend.mock_detector import detect
# 改为：
from inference.detector import detect
""")

h3('步骤 3.3 安装 ultralytics 并启动后端')
code_block("""# 在本机 venv（已装 CPU 版 torch）
cd "C:\\Users\\wzd\\Desktop\\毕业设计"
.\\.venv\\Scripts\\activate  # 或你当前用的 python 环境
pip install ultralytics
python src/backend/main.py
""")
p('服务启动后，浏览器访问 http://127.0.0.1:8000/docs 可看到 Swagger。')

h3('步骤 3.4 运行集成测试')
code_block("""# 另开一个终端
python scripts/test_m3.py
""")
p('预期 6/6 PASS。')

h2('第四阶段：W2/W3 同步提交里程碑汇报')
p('W2 填写并更新：W2_M2里程碑汇报_模型训练.docx，补充：')
p('实际训练耗时、最终 mAP 指标、分类模型准确率、Colab 取数是否顺利。')
p('W3 填写并更新：W3_M3里程碑汇报_后端可用.docx，补充：')
p('真实权重替换后的测试结果、API 文档链接、哈希链验证截图。')

h2('第五阶段：组长验收与启动 W4')
h3('验收 checklist')
rows = [
    ('检查项', '通过标准', '检查命令/位置'),
    ('W2 权重存在', 'models/best_detect.pt / best_cls.pt 存在', 'ls models'),
    ('W2 指标达标', 'mAP50-95 ≥ 0.90，分类准确率 ≥ 85%', 'W2_M2 汇报文档'),
    ('W3 接口跑通', 'scripts/test_m3.py 6/6 PASS', 'python scripts/test_m3.py'),
    ('Swagger 可访问', 'http://127.0.0.1:8000/docs 正常', '浏览器'),
    ('哈希链不可篡改', '篡改区块后 verify_chain() = False', 'scripts/test_m3.py 最后两项'),
]
table = doc.add_table(rows=1, cols=3)
table.style = 'Light Grid Accent 1'
for i, cell in enumerate(table.rows[0].cells):
    cell.text = rows[0][i]
for r in rows[1:]:
    row = table.add_row().cells
    for i, v in enumerate(r):
        row[i].text = v

h3('启动 W4 条件')
p('以上 5 项全部通过，即可组建 B 组（W4 前端工程师）。W4 将基于 W3 的 Swagger 文档开发 Vue 前端。')

h2('常见卡点与处理')
rows2 = [
    ('问题', '原因', '处理'),
    ('Colab 断开/超时', '免费 GPU 会话 12h 限制', '设置断点续训，权重定期 save 到 Drive'),
    ('Google Drive 上传慢', '网络或文件大', '压缩成 zip，晚上上传；或改用 Kaggle / 阿里云盘'),
    ('本地 CPU 推理慢', 'torch 是 CPU 版', '正常现象，只做验证；演示时可用 Colab 后端'),
    ('scripts/test_m3.py 报错', 'python-multipart 或 ultralytics 缺失', 'pip install python-multipart ultralytics'),
    ('detect() 返回字段不对', '没按契约②', '检查 detector.py / mock_detector.py 输出格式'),
]
table2 = doc.add_table(rows=1, cols=3)
table2.style = 'Light Grid Accent 1'
for i, cell in enumerate(table2.rows[0].cells):
    cell.text = rows2[0][i]
for r in rows2[1:]:
    row = table2.add_row().cells
    for i, v in enumerate(r):
        row[i].text = v

doc.save(OUT)
print(f"Saved: {OUT}")
