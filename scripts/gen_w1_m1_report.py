"""生成 W1 里程碑 M1 汇报文档（按派工指导手册汇报模板）。"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

OUT = "C:/Users/wzd/Desktop/毕业设计/W1_M1里程碑汇报_数据就绪.docx"

doc = Document()

# 全局中文字体
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(10.5)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def set_cn(run, font="宋体"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)


def h(text, level):
    p = doc.add_heading("", level=level)
    r = p.add_run(text)
    set_cn(r, "黑体")
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p


def para(text, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    set_cn(r)
    return p


# ===== 标题 =====
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("W1 里程碑 M1 汇报：数据集就绪")
r.bold = True
r.font.size = Pt(18)
set_cn(r, "黑体")

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("基于计算机视觉的农产品品质分级与溯源平台（恩施玉露茶）")
r.font.size = Pt(11)
set_cn(r, "楷体")

# ===== 汇报头信息表 =====
h("一、汇报头信息", 1)
info = [
    ("编号", "W1-M1-REPORT-001"),
    ("汇报人", "W1（数据工程工程师）"),
    ("日期", "2026-09-04"),
    ("里程碑", "M1（数据集就绪，计划 D3–D4，按期完成）"),
    ("关联变更", "CR-W1-001（数据集存储位置，已有条件批准）"),
]
t = doc.add_table(rows=len(info), cols=2)
t.style = "Table Grid"
t.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, (k, v) in enumerate(info):
    c0, c1 = t.rows[i].cells
    c0.width, c1.width = Cm(4), Cm(12)
    r0 = c0.paragraphs[0].add_run(k); r0.bold = True; set_cn(r0)
    r1 = c1.paragraphs[0].add_run(v); set_cn(r1)

# ===== 完成项 =====
h("二、完成项", 1)
items = [
    "数据集获取：TeaLeafAgeQuality（Mendeley Data，DOI 10.17632/7t964jmmy3，CC BY 4.0），全量 zip 4.59 GB，88 分片并行下载 + 断点续传，完整性校验通过（18104 条目可正常读取）。",
    "子集选择与提取：选用 Annotated 版本（2208 对图像/标注），弃用其 Augmented 版本（5298 对，与“W1 自行增强”要求冲突），弃用 Raw 无标注版本。13 张空标注背景图跳过（0.59%），实际入册 2195 对。",
    "类别映射确认：实测标注 class_id 0–3 与源数据 Category A/B/C/D 一一对应，映射为契约类别 0=特级 / 1=一级 / 2=二级 / 3=等外，id 无需重排。",
    "数据划分：按 7:2:1 分层随机划分（固定种子 42）——train 1537 / val 439 / test 219，各类别比例均衡。",
    "数据增强：仅对 train 集做 Albumentations 几何/光度增强（翻转、旋转、平移缩放、亮度对比度、HSV、高斯噪声/模糊），bbox 坐标同步变换，AUG_FACTOR=2，train 扩至 4611，val/test 不增强。",
    "契约交付：data.yaml（项目根 + 数据集内便携副本双份）、split_manifest.csv（逐样本 2195 行）、split_summary.csv、class_distribution.csv 均已生成。",
    "存储方案落地：按 CR-W1-001 批复，实际数据存 D:\\BISHE_DATA\\datasets\\tea_yulu，项目根 datasets/tea_yulu 以 Windows junction 映射，接口契约路径不变。",
    "文档：《数据说明.md》已回填全部真实数值（来源/映射/结构/预处理/规模/已知问题/Colab 取数方案/复现方式）。",
    "终检：scripts/w1_final_check.py 全量校验 22 项全部 PASS（存在性、junction 可达、契约符合性、图像标注一一配对、标注 0 异常行、清单行数）。",
]
for it in items:
    para("• " + it)

# ===== 交付物路径 =====
h("三、交付物路径（供组长核查）", 1)
deliverables = [
    ("数据集（经 junction 访问）", r"C:\Users\wzd\Desktop\毕业设计\datasets\tea_yulu", "train/val/test 各含 images 与 labels"),
    ("数据集实际存储位置", r"D:\BISHE_DATA\datasets\tea_yulu", "junction 目标，原始 zip 在 D:\\BISHE_DATA\\raw"),
    ("YOLO 配置（项目根）", r"C:\Users\wzd\Desktop\毕业设计\data.yaml", "path: datasets/tea_yulu，本地训练用"),
    ("YOLO 配置（便携副本）", r"D:\BISHE_DATA\datasets\tea_yulu\data.yaml", "path: .，供 Colab 整目录拷贝后使用"),
    ("逐样本划分清单", r"D:\BISHE_DATA\datasets\tea_yulu\split_manifest.csv", "2195 行：filename/split/class_id/class_name"),
    ("划分汇总统计", r"D:\BISHE_DATA\datasets\tea_yulu\split_summary.csv", "split × class 计数"),
    ("类别分布", r"D:\BISHE_DATA\datasets\tea_yulu\class_distribution.csv", "四类入册数量"),
    ("数据说明文档", r"C:\Users\wzd\Desktop\毕业设计\数据说明.md", "含 Colab 取数方案（第 7 节，供 W2）"),
    ("处理脚本", r"C:\Users\wzd\Desktop\毕业设计\scripts\w1_*.py、setup_data_link.bat", "下载/提取/划分/增强/建链/终检，种子 42 可复现"),
    ("终检脚本", r"C:\Users\wzd\Desktop\毕业设计\scripts\w1_final_check.py", "重跑即可复验本报告全部指标"),
]
t2 = doc.add_table(rows=len(deliverables) + 1, cols=3)
t2.style = "Table Grid"
hdr = t2.rows[0].cells
for j, htxt in enumerate(["交付物", "路径", "说明"]):
    rj = hdr[j].paragraphs[0].add_run(htxt); rj.bold = True; set_cn(rj)
for i, (a, b, c) in enumerate(deliverables, start=1):
    cells = t2.rows[i].cells
    for j, txt in enumerate([a, b, c]):
        rj = cells[j].paragraphs[0].add_run(txt); set_cn(rj)

# ===== 关键指标 =====
h("四、关键指标", 1)
metrics = [
    ("子集", "原始样本", "增强后样本", "说明"),
    ("train", "1537", "4611", "AUG_FACTOR=2，仅训练集增强"),
    ("val", "439", "439", "不增强"),
    ("test", "219", "219", "不增强"),
    ("合计", "2195", "5269", "跳过 13 张空标注背景图（0.59%）"),
]
t3 = doc.add_table(rows=len(metrics), cols=4)
t3.style = "Table Grid"
for i, row in enumerate(metrics):
    for j, txt in enumerate(row):
        rj = t3.rows[i].cells[j].paragraphs[0].add_run(txt)
        rj.bold = (i == 0)
        set_cn(rj)

para("")
para("bbox 类别分布（全量 5269 样本）：特级 1338 / 一级 1488 / 二级 1223 / 等外 1248，最大/最小比约 1.21，类别基本均衡，无需额外均衡处理。")
para("终检结果：22 项校验全部 PASS（含图像/标注一一对应、标注 0 异常行、data.yaml 契约① 符合、便携副本无盘符绝对路径）。")

# ===== 契约变更 =====
h("五、契约变更", 1)
para("无。类别 id、目录结构、data.yaml 字段均与接口契约①一致。存储位置变更（D 盘 + junction）已经 CR-W1-001 批准，不改变契约路径。", bold=False)

# ===== 阻塞风险 =====
h("六、阻塞风险", 1)
para("当前无阻塞。两项提示，不构成本里程碑阻塞：")
para("1. 13 张空标注背景图已跳过（0.59%）。若 W2 需要负样本参与检测训练，可决策后另行引入。")
para("2. 部分 bbox 接近整图范围，W2 可据此决策使用检测模型（YOLO 检测）或分类模型（yolov8-cls），需 W2 与组长确认。")

# ===== 下一步 =====
h("七、下一步", 1)
para("1. 组长核查本报告交付物（可重跑 scripts/w1_final_check.py 一键复验）。")
para("2. 启动 W2（视觉模型）：数据已就绪，Colab 取数方案见《数据说明.md》第 7 节（推荐方案 A：整目录上传 Google Drive 后挂载）。")
para("3. W1 转入待命：响应 W2 数据相关问题，必要时补充负样本或调整增强策略。")

# ===== 需组长决策项 =====
h("八、需组长决策项", 1)
para("1. 是否批准 W2 按《数据说明.md》第 7 节方案 A（Google Drive 上传）取数并开始训练？")
para("2. 检测模型 vs 分类模型的任务定位，请组长与 W2 确认后冻结（影响接口契约②的输出格式）。")

doc.save(OUT)
print(f"已生成: {OUT}")
