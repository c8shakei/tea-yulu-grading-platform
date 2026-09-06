"""Generate W3-M3 milestone report docx."""

import sys
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

OUTPUT = Path("c:/Users/wzd/Desktop/毕业设计/W3_M3里程碑汇报_后端可用.docx")


def set_cell(cell, text, bold=False):
    cell.text = str(text)
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.bold = bold
            run.font.size = Pt(10.5)


def main():
    doc = Document()

    title = doc.add_heading("W3-M3 里程碑汇报", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = doc.add_paragraph("基于计算机视觉的农产品品质分级与溯源平台（恩施玉露茶）")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()

    # 基本信息表格
    table = doc.add_table(rows=6, cols=2)
    table.style = "Light Grid Accent 1"
    fields = [
        ("编号", "W3-M3-REPORT-001"),
        ("日期", datetime.now().strftime("%Y-%m-%d")),
        ("里程碑", "M3 后端可用"),
        ("责任人", "W3 后端与溯源工程师"),
        ("作物", "恩施玉露茶"),
        ("技术栈", "YOLOv8 + 哈希链 + Vue + FastAPI + SQLite"),
    ]
    for i, (k, v) in enumerate(fields):
        set_cell(table.rows[i].cells[0], k, bold=True)
        set_cell(table.rows[i].cells[1], v)

    doc.add_heading("一、本阶段完成项", level=1)
    items = [
        "使用 FastAPI 搭建后端服务，统一响应体 {code, message, data}。",
        "实现接口契约③：POST /api/detect、POST /api/trace、GET /api/trace/{id}。",
        "设计 SQLite 单文件数据库，包含 detections、blocks 两张表并通过 trace_id 关联。",
        "实现 Python hashlib 哈希链模块（src/hashchain/chain.py），支持创世块、追加、完整性校验。",
        "提供 OpenAPI/Swagger 文档（/docs、/redoc、/openapi.json）。",
        "编写 mock 推理器（src/backend/mock_detector.py），W2 交付真实模型后仅替换该模块。",
        "完成 M3 集成测试：接口、OpenAPI、哈希链、篡改检测全部通过。",
    ]
    for item in items:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("二、交付物路径", level=1)
    paths = [
        "后端主入口：src/backend/main.py",
        "检测路由：src/backend/routers/detect.py",
        "溯源路由：src/backend/routers/trace.py",
        "数据库模块：src/backend/db.py",
        "哈希链模块：src/hashchain/chain.py",
        "Mock 推理：src/backend/mock_detector.py",
        "API 文档：docs/W3_M3_API文档.md",
        "数据库说明：docs/W3_M3_数据库说明.md",
        "测试脚本：scripts/test_m3.py",
        "数据库文件：data/tea_yulu.db",
        "汇报文档：W3_M3里程碑汇报_后端可用.docx",
    ]
    for p in paths:
        doc.add_paragraph(p, style="List Number")

    doc.add_heading("三、关键指标", level=1)
    metrics = [
        ("接口数量", "3 个业务接口 + 1 个健康检查"),
        ("数据库表", "2 张（detections、blocks）"),
        ("哈希链验证", "PASS（创世块+追加+完整性校验）"),
        ("篡改检测", "PASS（任一区块 data_hash 被篡改后 verify_chain() 返回 False）"),
        ("测试覆盖", "health、detect、trace POST、trace GET、openapi、tamper"),
        ("Swagger 文档", "可用：/docs、/openapi.json"),
    ]
    m_table = doc.add_table(rows=len(metrics), cols=2)
    m_table.style = "Light Grid Accent 1"
    for i, (k, v) in enumerate(metrics):
        set_cell(m_table.rows[i].cells[0], k, bold=True)
        set_cell(m_table.rows[i].cells[1], v)

    doc.add_heading("四、契约变更记录", level=1)
    doc.add_paragraph("无变更。严格按契约③、契约④实现；仅对 POST /api/trace 在 data 中额外返回 trace_id 以方便前端回显，不影响 W4 对接。")

    doc.add_heading("五、阻塞风险", level=1)
    risks = [
        "W2 真实 YOLOv8 推理权重交付时间：当前使用 mock 替代，接口签名已对齐，替换成本极低。",
        "W4 前端联调：本地 CORS 已开放，可直接对接。",
        "演示数据库文件 data/tea_yulu.db 与上传目录 uploads/ 已加入 .gitignore 建议（未强制），避免提交大文件。",
    ]
    for risk in risks:
        doc.add_paragraph(risk, style="List Bullet")

    doc.add_heading("六、下一步", level=1)
    next_steps = [
        "待 W2 交付真实 detect(image) 封装后，替换 src/backend/mock_detector.py。",
        "与 W4 前端完成 /api/detect、/api/trace、/api/trace/{id} 联调。",
        "补充异常场景单元测试（缺失 trace_id、空链查询、非法图片格式）。",
        "进入 M4 里程碑：前后端联调与演示页面。",
    ]
    for step in next_steps:
        doc.add_paragraph(step, style="List Bullet")

    doc.add_heading("七、需组长决策项", level=1)
    doc.add_paragraph("暂无。如遇 W2 接口签名与契约②不一致，将立即上报并申请变更。")

    doc.add_paragraph()
    sig = doc.add_paragraph("汇报人：W3 后端与溯源工程师")
    sig.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    doc.save(str(OUTPUT))
    print(f"报告已生成：{OUTPUT}")


if __name__ == "__main__":
    main()
