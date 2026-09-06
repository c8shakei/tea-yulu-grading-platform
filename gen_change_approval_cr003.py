# -*- coding: utf-8 -*-
"""生成 CR-W1-003 变更批复书 docx"""
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from datetime import datetime

OUT = "C:\\Users\\wzd\\Desktop\\毕业设计\\变更批复书_CR-W1-003.docx"

doc = Document()

style = doc.styles['Normal']
style.font.name = 'Microsoft YaHei'
style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
style.font.size = Pt(10.5)

title = doc.add_heading('变更批复书', level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

info_rows = [
    ('变更编号', 'CR-W1-003'),
    ('项目名称', '基于计算机视觉的农产品品质分级与溯源平台'),
    ('申请人', 'W1（数据工程工程师）'),
    ('关联里程碑', 'M2（模型训练）· yolov8s 追加实验 + 训练平台迁移'),
    ('申请日期', '2026-09-05'),
    ('批复日期', '2026-09-05'),
    ('批复人', '组长'),
]
table = doc.add_table(rows=1, cols=2)
table.style = 'Light Grid Accent 1'
table.rows[0].cells[0].text = '字段'
table.rows[0].cells[1].text = '内容'
for k, v in info_rows:
    row = table.add_row().cells
    row[0].text = k
    row[1].text = v

doc.add_paragraph()
doc.add_heading('一、变更内容摘要', level=2)
content = [
    '新增环境变量 WORK_DIR / SYNC_TARGET / DATA_ROOT，使 train_detect.py / train_cls.py / sync_checkpoints.py 同时支持 Colab + Google Drive 与 Kaggle 两种执行环境。',
    '数据集定位逻辑自适应：优先 WORK_DIR/datasets/tea_yulu，其次 Colab 本地目录，再回退 Drive zip，适配两平台。',
    'resolve_data_yaml() 改为写到可写目录，规避 Kaggle /kaggle/input 只读目录导致训练无法启动的问题。',
    'yolov8s 追加实验执行环境由 Colab 免费 T4 迁移至 Kaggle 免费 P100 GPU；训练脚本、数据、超参与产物命名规则均不变。',
]
for c in content:
    doc.add_paragraph(c, style='List Number')

doc.add_heading('二、批复结论', level=2)
doc.add_paragraph('同意方案 A：脚本平台可移植改造 + 迁移 Kaggle。')
doc.add_paragraph('附加说明：本次变更申请人字段填写为 W1，但变更对象属于 W2 视觉模型训练执行环境；结合项目经理确认，实际由 W2 发起并实施，W1 负责脚本工程化与回归校验。批复以文件编号 CR-W1-003 为准。')

doc.add_heading('三、关键决策', level=2)
doc.add_paragraph('① 训练平台迁移：')
doc.add_paragraph('批准将 yolov8s 追加实验由 Colab 免费 T4 迁移至 Kaggle 免费 P100。Colab 免费 GPU 滚动额度不可控，继续等待将阻塞 M2 进度；Kaggle P100 为零成本、性能更优的替代方案。')
doc.add_paragraph('② 接口契约②与产物命名：')
doc.add_paragraph('完全不变。无论 Colab 还是 Kaggle，产物文件名仍为 best_detect.pt / best_detect_yolov8s.pt / best_cls.pt 等；下载到本地 models\\ 后路径一致，W3 后端无需调整。')
doc.add_paragraph('③ 平台兼容性：')
doc.add_paragraph('默认环境变量仍走 Colab + Google Drive 旧流程；Kaggle 环境下显式设置 WORK_DIR=/kaggle/working、SYNC_TARGET=/kaggle/working/BISHE/models 即可。两种路径均须通过实测验证。')

doc.add_heading('四、生效范围与后续动作', level=2)
actions = [
    'W2 在 Kaggle 新建 Notebook，开启 GPU（P100），上传 tea_yulu_dataset.zip 与三脚本（train_detect.py / train_cls.py / sync_checkpoints.py）。',
    'Kaggle 执行前设置环境变量：%env WORK_DIR=/kaggle/working；%env SYNC_TARGET=/kaggle/working/BISHE/models；%env YOLO_MODEL_SIZE=yolov8s。',
    '解压数据集到 /kaggle/working/datasets/tea_yulu，确认 data.yaml 存在后运行训练脚本。',
    '训练完成后从 Kaggle 输出区下载 best_detect_yolov8s.pt / best_cls_yolov8s-cls.pt 等到本机 models\\ 目录。',
    'Colab 旧流程与 yolov8n 基线保持不变；未来 Colab 额度恢复后仍可复用。',
]
for a in actions:
    doc.add_paragraph(a, style='List Number')

doc.add_heading('五、风险与约束', level=2)
risks = [
    'Kaggle 免费 GPU 额度约 30h/周滚动，yolov8s 检测+分类训练预计可在一个窗口内完成，但仍需关注会话稳定性。',
    'Kaggle /kaggle/input 为只读目录，所有写入（包括 resolve_data_yaml 生成的 data_colab.yaml）必须落到 /kaggle/working。',
    'Kaggle Notebook 默认不持久化 /kaggle/working 之外的路径，训练产物必须通过 SYNC_TARGET 落到 /kaggle/working 后再下载。',
    '迁移后不得改动类别映射、类别顺序、接口契约②字段；任何训练超参调整仍需在变更申请中说明。',
]
for r in risks:
    doc.add_paragraph(r, style='List Bullet')

doc.add_paragraph()
p = doc.add_paragraph('批复人签字：_______________')
p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
p = doc.add_paragraph('日期：2026-09-05')
p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

doc.save(OUT)
print(f"Saved: {OUT}")
