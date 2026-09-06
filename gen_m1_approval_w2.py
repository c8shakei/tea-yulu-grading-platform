#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
M1 验收批复书暨 W2 启动令生成器
"""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from datetime import datetime
import os

OUTPUT = r"C:\Users\wzd\Desktop\毕业设计\M1验收批复书_W2启动令.docx"

def set_default_font(doc, name="宋体", size=10.5):
    doc.styles["Normal"].font.name = name
    doc.styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    doc.styles["Normal"].font.size = Pt(size)

def add_heading_zh(doc, text, level=1):
    style = "Heading %d" % level
    p = doc.add_paragraph(style=style)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.font.name = "黑体" if level <= 2 else "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体" if level <= 2 else "宋体")
    run.font.bold = True
    run.font.size = Pt({1: 16, 2: 14}.get(level, 12))
    return p

def add_para(doc, text, bold=False, indent=True, size=10.5):
    p = doc.add_paragraph(style="Normal")
    if indent:
        p.paragraph_format.first_line_indent = Inches(0.3)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    run = p.add_run(text)
    run.font.bold = bold
    run.font.size = Pt(size)
    return p

def add_table(doc, rows, cols, data):
    table = doc.add_table(rows=rows, cols=cols)
    table.style = "Table Grid"
    for i, row_data in enumerate(data):
        row = table.rows[i]
        for j, cell_text in enumerate(row_data):
            cell = row.cells[j]
            cell.text = str(cell_text)
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    return table

def main():
    doc = Document()
    set_default_font(doc)

    # Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("M1 里程碑验收批复书\n暨 W2 启动令")
    run.font.name = "黑体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    run.font.size = Pt(22)
    run.font.bold = True

    # Subtitle
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run("项目：基于计算机视觉的农产品品质分级与溯源平台（恩施玉露茶）")
    run.font.size = Pt(12)
    run = sub.add_run("\n批复编号：APPROVAL-M1-W2-20260904    批复日期：2026-09-04")
    run.font.size = Pt(12)

    add_heading_zh(doc, "一、基本信息", level=1)
    info = [
        ["项", "内容"],
        ["汇报编号", "W1-M1-REPORT-001"],
        ["汇报人", "W1（数据工程工程师）"],
        ["验收里程碑", "M1 数据集就绪（计划 D3–D4）"],
        ["批复人", "项目组长"],
        ["关联变更", "CR-W1-001（数据集存储位置，已有条件批准）"],
        ["待启动下游", "W2 视觉模型工程师（M2 模型训练与封装）"],
    ]
    add_table(doc, len(info), 2, info)

    add_heading_zh(doc, "二、M1 验收结论", level=1)
    add_para(doc, "经审阅《W1_M1里程碑汇报_数据就绪.docx》、复核《数据说明.md》关键章节，并复跑 scripts/w1_final_check.py 进行全量校验，结论如下：")
    add_para(doc, "M1 里程碑验收通过。数据集已按接口契约①完成获取、类别映射、划分、增强与存储方案落地，满足全部验收口径，具备启动 W2 的条件。", bold=True)

    add_heading_zh(doc, "三、验收口径核对表", level=1)
    check = [
        ["验收项", "标准", "W1 交付", "复核结果"],
        ["junction 可达", "项目根 datasets/tea_yulu 能正常访问 D 盘数据", "已建 junction：C:\\...\\datasets\\tea_yulu → D:\\BISHE_DATA\\datasets\\tea_yulu", "PASS"],
        ["data.yaml 无盘符绝对路径", "data.yaml 中不得出现 D:\\ 或 C:\\ 等绝对路径", "项目根 data.yaml：path: datasets/tea_yulu；便携副本：path: .", "PASS"],
        ["划分可复现", "固定随机种子，重复执行结果一致", "种子 42，7:2:1 分层划分，split_manifest.csv 2195 行", "PASS"],
        ["Colab 取数方案", "《数据说明》须写明 W2 在 Colab 如何取数", "第 7 节含方案 A（Google Drive 上传挂载）与方案 B（Mendeley 重下）", "PASS"],
        ["终检脚本", "w1_final_check.py 全部 PASS", "22/22 项全部 PASS", "PASS"],
    ]
    add_table(doc, len(check), 4, check)

    add_heading_zh(doc, "四、组长决策批复", level=1)
    add_heading_zh(doc, "4.1 Colab 取数方案", level=2)
    add_para(doc, "批准方案 A（Google Drive 上传挂载）为主方案：将 D:\\BISHE_DATA\\datasets\\tea_yulu 压缩上传至 Google Drive，Colab 挂载后软链到 /content/datasets/tea_yulu，使用便携版 data.yaml（path: .）启动训练。")
    add_para(doc, "方案 B（在 Colab 内通过 w1_download_dataset.py 与 w1_split_dataset.py 重新下载准备）作为 Fallback。")

    add_heading_zh(doc, "4.2 检测 vs 分类任务定位", level=2)
    add_para(doc, "主模型维持 YOLOv8 检测任务（yolov8n 起步，可选 yolov8s），继续按接口契约②输出 {class_id, class_name, confidence, bbox}。W3 后端、W4 前端、W5 集成测试均以检测接口为默认消费路径。")
    add_para(doc, "W2 须额外训练一个 yolov8n-cls 分类模型作为对照实验，记录分类准确率，与检测模型 mAP 等指标一并写入论文对比章节。该对照模型不影响 W3/W4 默认集成路径。")

    add_heading_zh(doc, "五、非阻塞改进项", level=1)
    add_para(doc, "建议 W1 在待命期间把 13 张空标注背景图作为负样本补入 train 集：复制图像到 train/images，并在 train/labels 中放置同名空 .txt 文件，更新 split_manifest.csv 与 w1_final_check.py 计数。")
    add_para(doc, "该改进不阻塞 W2 启动，W2 可先用现有 4611 张训练集开始训练；W1 在并行待命时顺手补齐，以提升检测模型对无目标场景的鲁棒性。")

    add_heading_zh(doc, "六、W2 启动令", level=1)
    add_para(doc, "自本批复签署之日起，W2 视觉模型工程师正式承担 M2 里程碑任务：")
    add_heading_zh(doc, "6.1 任务目标", level=2)
    add_para(doc, "1. 在 Colab GPU 上完成 YOLOv8 检测模型训练，产出 best.pt；")
    add_para(doc, "2. 训练 yolov8n-cls 分类模型作为对照实验；")
    add_para(doc, "3. 封装 detect(image) 推理接口，按契约②输出标准 JSON；")
    add_para(doc, "4. 输出训练曲线、混淆矩阵、mAP/准确率对比表及模型说明文档。")

    add_heading_zh(doc, "6.2 交付物", level=2)
    deliver = [
        ["交付物", "路径/命名建议", "说明"],
        ["检测模型权重", "models/best_detect.pt", "YOLOv8 检测主模型"],
        ["分类对照模型权重", "models/best_cls.pt", "yolov8n-cls 对照"],
        ["推理封装脚本", "src/inference/detector.py", "detect(image) → {class_id, class_name, confidence, bbox}"],
        ["训练说明", "docs/W2_M2_训练说明.md", "超参数、训练时长、Colab 环境、复现方式"],
        ["对比实验结果", "docs/W2_M2_对比实验.md", "检测 mAP vs 分类准确率、混淆矩阵"],
        ["M2 里程碑汇报", "W2_M2里程碑汇报_模型训练.docx", "按汇报模板填写"],
    ]
    add_table(doc, len(deliver), 3, deliver)

    add_heading_zh(doc, "6.3 关键约束", level=2)
    add_para(doc, "• 不得改动 data.yaml 的类别顺序、类别数量、相对路径结构；如需调整须先报组长并同步 W1/W3。")
    add_para(doc, "• detect() 接口输出字段必须与接口契约②一致：class_id, class_name, confidence, bbox（[x_center, y_center, width, height] 归一化 0–1）。")
    add_para(doc, "• 训练过程中定期保存 checkpoint 到 Google Drive，防止 Colab 会话中断导致训练成果丢失。")
    add_para(doc, "• 里程碑 M2 计划 D10（2026-09-14 前）完成；若遇阻塞须提前 24 小时上报。")

    add_heading_zh(doc, "七、生效与签署", level=1)
    sign = [
        ["角色", "签署", "日期"],
        ["申请人 / W1", "______________", "______________"],
        ["批复人 / 组长", "（已电子批复）", "2026-09-04"],
        ["承接人 / W2", "______________", "______________"],
    ]
    add_table(doc, len(sign), 3, sign)

    add_para(doc, "注：本批复书作为项目变更控制与里程碑管理的过程文档，与《项目章程与进度计划》《软件需求规格说明书 SRS》《接口契约①/②》配套存档。")

    doc.save(OUTPUT)
    print("OK saved:", OUTPUT)

if __name__ == "__main__":
    main()
