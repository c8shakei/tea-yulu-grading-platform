# -*- coding: utf-8 -*-
"""生成 CR-W1-002 变更批复书 docx"""
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from datetime import datetime

OUT = "C:\\Users\\wzd\\Desktop\\毕业设计\\变更批复书_CR-W1-002.docx"

doc = Document()

style = doc.styles['Normal']
style.font.name = 'Microsoft YaHei'
style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
style.font.size = Pt(10.5)

title = doc.add_heading('变更批复书', level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

info_rows = [
    ('变更编号', 'CR-W1-002'),
    ('项目名称', '基于计算机视觉的农产品品质分级与溯源平台'),
    ('申请人', 'W1（数据工程工程师）'),
    ('关联里程碑', 'M2（模型训练）· yolov8s 追加实验'),
    ('申请日期', '2026-09-05'),
    ('批复日期', '2026-09-05'),
    ('批复人', '组长'),
]
table = doc.add_table(rows=1, cols=2)
table.style = 'Light Grid Accent 1'
table.rows[0].cells[0].text = '字段'
table.rows[0].cells[1].text = '内容'
for k, v in info_rows:
    row = table.add_row().cells
    row[0].text = k
    row[1].text = v

doc.add_paragraph()
doc.add_heading('一、变更内容摘要', level=2)
content = [
    '1. train_detect.py / train_cls.py 的模型规模由硬编码 yolov8n 改为通过环境变量 YOLO_MODEL_SIZE 可配置；未设置时默认 yolov8n，保持旧行为完全兼容。',
    '2. 训练产物同步到 Google Drive 时文件名追加规模后缀（如 best_detect_yolov8s.pt），实现 yolov8n 基线与 yolov8s 实验的物理隔离。',
    '3. yolov8s 规模下自动下调 batch size：检测 16→8，分类 64→32，规避 T4 GPU 显存溢出。',
    '4. sync_checkpoints.py 同步支持 YOLO_MODEL_SIZE，收尾崩溃时可就地抢救对应规模产物。',
]
for c in content:
    doc.add_paragraph(c, style='List Number')

doc.add_heading('二、批复结论', level=2)
doc.add_paragraph('同意方案 A：环境变量可配 + 文件名加后缀物理隔离。')
doc.add_paragraph('附加说明：变更申请人字段虽填写为 W1，但变更对象（训练脚本）属于 M2 视觉模型工作域；结合项目经理确认，本次变更实际由 W2 视觉模型工程师发起并实施，W1 承担脚本工程化与回归校验。批复以文件编号 CR-W1-002 为准。')

doc.add_heading('三、关键决策', level=2)
doc.add_paragraph('① 论文主模型是否升为 yolov8s：')
doc.add_paragraph('暂缓决策。当前维持 yolov8n 为论文主模型与系统部署默认模型；yolov8s 仅作为消融对照实验。待 yolov8s 训练完成后，按以下量化口径二次裁定：')
criteria = [
    '若 yolov8s 相比 yolov8n 的 mAP50-95 提升 ≥ 0.05，且本机 CPU 推理单张耗时 < 200ms，则批准将主模型切换为 yolov8s；',
    '否则维持 yolov8n 为主模型，yolov8s 仅用于论文消融对比章节。',
]
for c in criteria:
    doc.add_paragraph(c, style='List Bullet')

doc.add_paragraph('② 接口契约②权重路径：')
doc.add_paragraph('在二次裁定前，W3 后端仍按原契约加载 best_detect.pt（yolov8n）。若主模型切换为 yolov8s，须由 W2 另提变更申请，同步修改接口契约②与 W3 加载逻辑。')

doc.add_heading('四、生效范围与后续动作', level=2)
actions = [
    'W2 可在 Colab 立即执行 yolov8s 追加实验：先 %env YOLO_MODEL_SIZE=yolov8s，再 %run train_detect.py / train_cls.py。',
    'yolov8n 基线产物（best_detect.pt 等）保留在 Drive BISHE/models/，禁止覆盖。',
    'W3 后端保持现状，继续以 best_detect.pt（yolov8n）对接 mock→真实模型切换工作。',
    'yolov8s 实验完成后，W2 须在 M2 汇报中单独列出 yolov8n vs yolov8s 消融指标，并提出主模型切换的后续变更申请（如适用）。',
]
for a in actions:
    doc.add_paragraph(a, style='List Number')

doc.add_heading('五、风险与约束', level=2)
risks = [
    'yolov8s 训练时长约为 yolov8n 的 2–3 倍，Colab 免费 GPU 有 12h 会话限制，需设置断点续训或分阶段运行。',
    '切换主模型属于接口契约②变更，必须走正式变更流程，不得由 W2/W3 私下协商。',
    '所有训练产物命名须严格按 YOLO_MODEL_SIZE 后缀规则，避免 Drive 文件互相覆盖导致对照组丢失。',
]
for r in risks:
    doc.add_paragraph(r, style='List Bullet')

doc.add_paragraph()
p = doc.add_paragraph('批复人签字：_______________')
p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
p = doc.add_paragraph('日期：2026-09-05')
p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

doc.save(OUT)
print(f"Saved: {OUT}")
